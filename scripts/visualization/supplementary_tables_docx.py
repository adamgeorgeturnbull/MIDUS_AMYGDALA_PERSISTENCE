#!/usr/bin/env python3
"""
supplementary_tables_docx.py

Reads results/tables/supplementary/supp_all_methods.csv and writes a
formatted Word (.docx) table suitable for submission as a supplementary table.

Formatting:
  - Header row: bold, dark grey background, white text
  - Analysis groups: light grey shading on first row of each new analysis
  - Alternating white / very light grey rows within each analysis group
  - Significant p-values (containing *) rendered in bold
  - Landscape page, Times New Roman 10pt, 1 inch margins

Run from project root directory.
"""

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT

CSV_PATH = Path("results/tables/supplementary/supp_all_methods.csv")
OUT_PATH = Path("results/tables/supplementary/supp_all_methods.docx")

# ── Colours ───────────────────────────────────────────────────────────────────
HEADER_BG   = "2F2F2F"   # dark grey header
GROUP_BG    = "D9D9D9"   # light grey — first row of each analysis group
ALT_BG      = "F2F2F2"   # very light grey — alternating rows
WHITE       = "FFFFFF"

FONT_NAME   = "Times New Roman"
FONT_SIZE   = Pt(10)
HEADER_SIZE = Pt(10)

# Columns that are right-aligned (numeric)
RIGHT_COLS = {"N", "r", "r p", "OLS t", "OLS p", "MLM z", "MLM p"}


def set_cell_bg(cell, hex_color):
    shading = parse_xml(
        f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hex_color}"/>'
    )
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_text(cell, text, bold=False, font_name=FONT_NAME,
                  font_size=FONT_SIZE, color=None,
                  align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = align
    run = para.add_run(str(text))
    run.bold = bold
    run.font.name = font_name
    run.font.size = font_size
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def main():
    df = pd.read_csv(CSV_PATH)

    doc = Document()

    # ── Page layout: landscape ────────────────────────────────────────────────
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    margin = Inches(1.0)
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, attr, margin)

    # ── Title ─────────────────────────────────────────────────────────────────
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run(
        "Supplementary Table. Correlation, OLS regression, and MLM results "
        "for primary analyses (Analyses 01-05)."
    )
    run.bold = True
    run.font.name = FONT_NAME
    run.font.size = Pt(10)

    note = doc.add_paragraph()
    note_run = note.add_run(
        "p-values are one-tailed for directional hypotheses (Analyses 01-05) "
        "and two-tailed otherwise. * p < .05, ** p < .01, *** p < .001."
    )
    note_run.font.name = FONT_NAME
    note_run.font.size = Pt(9)
    note_run.italic = True

    # ── Table ─────────────────────────────────────────────────────────────────
    cols = list(df.columns)
    n_cols = len(cols)

    table = doc.add_table(rows=1, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # Set column widths
    # Analysis ~2.2in, Predictor ~1.6in, Outcome ~1.4in, N ~0.4in, rest ~0.7in each
    col_widths = {
        "Analysis":  Inches(2.2),
        "Predictor": Inches(1.6),
        "Outcome":   Inches(1.4),
        "N":         Inches(0.4),
    }
    default_w = Inches(0.75)

    for i, col in enumerate(cols):
        for cell in table.columns[i].cells:
            cell.width = col_widths.get(col, default_w)

    # Header row
    hdr = table.rows[0]
    for i, col in enumerate(cols):
        cell = hdr.cells[i]
        set_cell_bg(cell, HEADER_BG)
        align = WD_ALIGN_PARAGRAPH.RIGHT if col in RIGHT_COLS else WD_ALIGN_PARAGRAPH.LEFT
        set_cell_text(cell, col, bold=True, color=WHITE, align=align)

    # Data rows
    prev_analysis = None
    row_in_group  = 0

    for _, data_row in df.iterrows():
        analysis = data_row["Analysis"]
        is_new_group = analysis != prev_analysis

        if is_new_group:
            row_in_group = 0
            prev_analysis = analysis

        bg = GROUP_BG if is_new_group else (ALT_BG if row_in_group % 2 == 1 else WHITE)
        row_in_group += 1

        tr = table.add_row()
        for i, col in enumerate(cols):
            cell = tr.cells[i]
            set_cell_bg(cell, bg)
            val  = "" if pd.isna(data_row[col]) else str(data_row[col])
            bold = "*" in val and col in RIGHT_COLS  # bold significant p-values
            align = WD_ALIGN_PARAGRAPH.RIGHT if col in RIGHT_COLS else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cell, val, bold=bold, align=align)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    print(f"Saved: {OUT_PATH}")


if __name__ == "__main__":
    main()
